from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "flyer_interview_lab/test_materials/synthetic_data_analytics_resume.docx"
RED = RGBColor(206, 17, 65)
DARK = RGBColor(23, 32, 51)
GRAY = RGBColor(82, 98, 121)


def set_font(run, name="Aptos", size=10, bold=False, color=DARK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text.upper())
    set_font(r, size=10, bold=True, color=RED)
    bottom = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "6")
    border.set(qn("w:color"), "D9D9D9")
    bottom.append(border)
    p._p.get_or_add_pPr().append(bottom)


def add_entry(doc, title, detail, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    left = p.add_run(title)
    set_font(left, bold=True)
    right = p.add_run(f"  |  {detail}")
    set_font(right, color=GRAY)
    for text in bullets:
        b = doc.add_paragraph(style="List Bullet")
        b.paragraph_format.left_indent = Inches(0.2)
        b.paragraph_format.first_line_indent = Inches(-0.15)
        b.paragraph_format.space_after = Pt(1.5)
        set_font(b.add_run(text), size=9.5)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.45)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(2)

# Some Word installations give the built-in Title style a colored bottom rule.
# Keep the semantic Title style while explicitly removing that inherited border.
title_style_ppr = doc.styles["Title"].element.get_or_add_pPr()
for border in title_style_ppr.findall(qn("w:pBdr")):
    title_style_ppr.remove(border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(2)
r = title.add_run("Jordan Lee Synthetic Test Resume")
set_font(r, size=20, bold=True, color=DARK)

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(5)
set_font(
    contact.add_run(
        "jordan.lee@example.com  |  937-555-0142  |  123 Test Avenue, Dayton, OH 45469"
    ),
    size=9,
    color=GRAY,
)

notice = doc.add_paragraph()
notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
notice.paragraph_format.space_after = Pt(4)
set_font(
    notice.add_run("SYNTHETIC TEST DATA - NOT A REAL PERSON"),
    size=8.5,
    bold=True,
    color=RED,
)

add_section_heading(doc, "Education")
add_entry(
    doc,
    "University of Dayton",
    "Bachelor of Science in Data Analytics, expected May 2027 | GPA 3.48",
    [
        "Relevant coursework: Applied Statistics, Regression Analysis, Database Management, Data Visualization, and Business Analytics.",
        "Minor in Mathematics; Dean's List for two semesters.",
    ],
)

add_section_heading(doc, "Technical Skills")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
set_font(p.add_run("Programming and analysis: "), bold=True)
set_font(p.add_run("Python, R, Excel, introductory SQL, pandas, NumPy, ggplot2"))
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
set_font(p.add_run("Methods: "), bold=True)
set_font(p.add_run("data cleaning, regression, hypothesis testing, dashboard design, descriptive statistics"))

add_section_heading(doc, "Projects")
add_entry(
    doc,
    "Student Retention Analysis",
    "Applied Statistics course project | Spring 2026",
    [
        "Cleaned and joined three anonymized datasets containing 18,000 student records using Python and pandas.",
        "Built logistic regression models to explore factors associated with first-year retention and compared performance using cross-validation.",
        "Presented limitations and actionable findings in a ten-minute team presentation to faculty and classmates.",
    ],
)
add_entry(
    doc,
    "Regional Employment Dashboard",
    "Data Visualization course | Fall 2025",
    [
        "Organized public labor-market data in Excel and created an interactive dashboard tracking employment and wage trends by occupation.",
        "Wrote a two-page executive summary translating the analysis for a nontechnical audience.",
    ],
)

add_section_heading(doc, "Experience and Leadership")
add_entry(
    doc,
    "Mathematics Tutor",
    "University of Dayton Learning Teaching Center | August 2025 to present",
    [
        "Explain statistics and calculus concepts to 8-12 students per week and adapt explanations to different levels of preparation.",
        "Maintain concise session notes and coordinate with other tutors during high-demand exam periods.",
    ],
)
add_entry(
    doc,
    "Treasurer",
    "Data Science Club | January 2026 to present",
    [
        "Track a $2,500 annual budget in Excel and provide monthly updates to the student leadership team.",
        "Helped organize two alumni career panels attended by more than 60 students.",
    ],
)

doc.save(OUTPUT)
